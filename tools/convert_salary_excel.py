#!/usr/bin/env python3
"""Turn a salary spreadsheet into the JSON that salary_lookup.py reads.

Pay surveys arrive as spreadsheets, one per publisher, and no two of them agree
on anything: the header row is rarely the first row, the company column is
called whatever the publishing body calls a company in its own language, and
the figures come either as one column per category or as a count column paired
with an index column. This script reads that shape rather than requiring one,
and writes data/profile/salary_data.json beside the rest of your own data.

    python tools/convert_salary_excel.py survey.xlsx
    python tools/convert_salary_excel.py survey.xlsx --source "Union survey 2026"
    python tools/convert_salary_excel.py survey.xlsx --baseline 100 \
        --baseline-desc "100 = the median across all respondents"

What the sheet has to have: a header row somewhere in the first ten, a column
naming the company, and some numeric columns. A city column is used when it is
there. Everything else is inferred.

The one dependency in this repository lives here: openpyxl, imported inside
main() so that --help still answers on a machine that never installed it.
"""

import json
import re
import sys
import argparse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import paths  # noqa: E402


# What a header can call each thing, across the languages these surveys are
# published in. Add to these freely -- an unrecognised header costs you a
# column, and a wrong guess costs you a wrong benchmark.
COMPANY_PATTERNS = {
    # English
    "company", "employer", "organisation", "organization",
    # Danish and Norwegian
    "firma", "virksomhed", "arbejdsgiver",
    # Portuguese and Spanish
    "empresa", "empregador", "empleador", "companhia", "entidade",
    # Russian
    "компания", "работодатель", "организация",
}
CITY_PATTERNS = {
    # English
    "city", "location", "site",
    # Danish and Norwegian
    "by", "kommune", "lokation", "sted",
    # Portuguese and Spanish
    "cidade", "ciudad", "localidade", "localidad", "municipio", "concelho",
    # Russian
    "город", "местоположение",
}
COUNT_PATTERNS = {
    # English
    "count", "number", "n", "employees", "headcount", "respondents",
    # Danish and Norwegian
    "antal", "medarbejdere",
    # Portuguese and Spanish
    "quantidade", "cantidad", "empregados", "empleados", "trabalhadores",
    "respondentes",
    # Russian
    "количество", "число", "сотрудников",
}
INDEX_PATTERNS = {
    # English
    "index", "idx", "salary", "median", "average", "pay",
    # Danish and Norwegian
    "indeks", "løn", "gennemsnit",
    # Portuguese and Spanish
    "índice", "indice", "salário", "salario", "mediana", "média", "promedio",
    "remuneração", "remuneración", "vencimento",
    # Russian
    "зарплата", "оклад", "медиана", "индекс",
}


# Below this length a marker is trusted only as a whole word. "n" is a real
# header for a count on its own, and a substring of "Indeks", "Median" and
# "Engineering" alike -- matched as a substring it made every column a count,
# and the count/index pairing below then had nothing left to pair.
MIN_SUBSTRING_MARKER = 3


def header_words(header):
    """The header as lowercase words, with non-ASCII letters left alone."""
    return re.findall(r"[^\W_]+", (header or "").lower(), flags=re.UNICODE)


def marker_in(word):
    """Where a count or index marker ends inside one word, if it is in there.

    Returns (end offset, kind) for whichever marker ends last, or None. Danish
    and German weld these onto the category -- "Lønindeks",
    "Medarbejderantal" -- and such compounds are head-final, so the marker that
    ends last is the one saying what the column actually holds.
    """
    hits = [
        (word.rindex(p) + len(p), kind)
        for kind, patterns in (("count", COUNT_PATTERNS), ("index", INDEX_PATTERNS))
        for p in patterns
        if len(p) >= MIN_SUBSTRING_MARKER and p in word
    ]
    return max(hits) if hits else None


def detect_column_type(header):
    """Whether a header names a headcount, a salary index, or neither.

    A whole word settles it outright; substrings are consulted only afterwards,
    and only for markers long enough to be trusted that way. Whole words alone
    call "Lønindeks" and "Medarbejderantal" neither -- and an unclassified
    column is filed below as an index, which is how a headcount once came out
    of salary_lookup dressed as a benchmark figure.
    """
    words = header_words(header)
    if set(words) & COUNT_PATTERNS:
        return "count"
    if set(words) & INDEX_PATTERNS:
        return "index"
    hits = [m for m in (marker_in(w) for w in words) if m]
    return hits[-1][1] if hits else None


def category_name(header, patterns):
    """The header with its marker words taken out, as a snake_case name.

    Only words that actually matched are dropped, so "Antal Engineering" comes
    back as "engineering". Stripping the patterns as substrings instead took
    their letters out of the remaining words -- "atal egieerig" -- and did it in
    set-iteration order, so two runs disagreed. A word that merely contains a
    marker goes as well, or a compound keeps it: "lønindeks_engineering".
    """
    return "_".join(
        w for w in header_words(header)
        if w not in patterns
        and not any(len(p) >= MIN_SUBSTRING_MARKER and p in w for p in patterns)
    )


def find_header_row(ws):
    """The first row in the top ten holding a cell that names the company."""
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=False), start=1):
        for cell in row:
            if cell.value and str(cell.value).strip().lower() in COMPANY_PATTERNS:
                return row_idx
    return None


def read_headers(ws, header_row):
    """The header row as a list of stripped strings, blanks included."""
    return [str(cell.value).strip() if cell.value else "" for cell in ws[header_row]]


def locate_key_columns(headers):
    """Indices of the company and city columns, either of which may be absent.

    Matched on the whole header only. A substring match here would claim
    "Company car" as the company column.
    """
    company_col = None
    city_col = None
    for i, header in enumerate(headers):
        lowered = header.lower()
        if lowered in COMPANY_PATTERNS:
            company_col = i
        elif lowered in CITY_PATTERNS:
            city_col = i
    return company_col, city_col


def group_categories(data_cols):
    """Fold the data columns into categories, pairing counts with indices.

    A count column followed by an index column (or the reverse) describes one
    category between them, and the category's name is whatever is left of the
    header once the marker is removed. Anything unpaired stands alone under its
    own header.

    Returns (categories, unclassified) -- the second being the headers that
    could not be read as either kind, which the caller reports. They are still
    stored as an index, because there is nowhere else to put them, and that is
    exactly why they are worth naming out loud.
    """
    categories = []
    unclassified = []

    i = 0
    while i < len(data_cols):
        col_idx, header = data_cols[i]
        kind = detect_column_type(header)

        if i + 1 < len(data_cols):
            next_idx, next_header = data_cols[i + 1]
            next_kind = detect_column_type(next_header)

            pairing = None
            if kind == "count" and next_kind == "index":
                pairing = {"count_col": col_idx, "index_col": next_idx}
                strip = COUNT_PATTERNS
            elif kind == "index" and next_kind == "count":
                pairing = {"index_col": col_idx, "count_col": next_idx}
                strip = INDEX_PATTERNS

            if pairing is not None:
                name = category_name(header, strip) or f"category_{len(categories) + 1}"
                categories.append({"name": name, **pairing})
                i += 2
                continue

        if kind is None:
            unclassified.append(header)
        categories.append({
            "name": header.lower().replace(" ", "_"),
            "value_col": col_idx,
            "value_type": kind or "index",
        })
        i += 1

    return categories, unclassified


def cell_number(row, col, cast):
    """One cell read as a number, or None when it is empty or is not one."""
    if col is None or col >= len(row) or row[col] is None:
        return None
    try:
        return cast(row[col])
    except (ValueError, TypeError):
        return None


def read_company_rows(ws, header_row, company_col, city_col, categories):
    """Every row below the header, as one entry per named company."""
    entries = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row[company_col]:
            continue

        entry = {
            "company": str(row[company_col]).strip(),
            "city": str(row[city_col]).strip() if city_col is not None and row[city_col] else "",
            "categories": {},
        }

        for cat in categories:
            if "count_col" in cat and "index_col" in cat:
                entry["categories"][cat["name"]] = {
                    "count": cell_number(row, cat["count_col"], int),
                    "index": cell_number(row, cat["index_col"], float),
                }
            elif "value_col" in cat:
                col = cat["value_col"]
                if col < len(row) and row[col] is not None:
                    value = row[col]
                    try:
                        value = float(value)
                    except (ValueError, TypeError):
                        value = str(value)
                    entry["categories"][cat["name"]] = {cat["value_type"]: value}

        entries.append(entry)
    return entries


def parse_sheet(ws):
    """One worksheet as a list of company entries. Empty when it is unreadable."""
    header_row = find_header_row(ws)
    if header_row is None:
        print(
            f"Warning: no header row found in the first ten rows of sheet "
            f"'{ws.title}'. Skipping it.",
            file=sys.stderr,
        )
        return []

    headers = read_headers(ws, header_row)
    company_col, city_col = locate_key_columns(headers)
    if company_col is None:
        print(
            f"Warning: sheet '{ws.title}' has a header row but no column naming "
            "the company. Skipping it.",
            file=sys.stderr,
        )
        return []

    data_cols = [
        (i, header) for i, header in enumerate(headers)
        if header and i != company_col and i != city_col
    ]
    categories, unclassified = group_categories(data_cols)

    if unclassified:
        print(
            f"Warning: in sheet '{ws.title}', these columns could be a headcount "
            "or a salary figure and the header does not say which. Stored as a "
            "salary index: " + ", ".join(unclassified),
            file=sys.stderr,
        )

    return read_company_rows(ws, header_row, company_col, city_col, categories)


def main():
    parser = argparse.ArgumentParser(
        description="Convert a salary spreadsheet into salary_lookup's JSON",
    )
    parser.add_argument("excel_file", help="The spreadsheet to read")
    parser.add_argument(
        "--output", default=None,
        help="Where to write it (default: data/profile/salary_data.json)",
    )
    parser.add_argument(
        "--source", default=None,
        help="Who published the figures, recorded so a benchmark can cite them "
             "(default: the file's own name)",
    )
    parser.add_argument(
        "--baseline", type=float, default=100,
        help="What the index is measured against (default: 100)",
    )
    parser.add_argument(
        "--baseline-desc", default=None,
        help="What that baseline means in words, e.g. '100 = the median'",
    )
    args = parser.parse_args()

    # Imported here, not at module level, so --help answers on a machine that
    # never installed the one optional dependency in this repository.
    try:
        import openpyxl
    except ImportError:
        print(
            "Error: this is the one script here that needs a package. "
            "Install it with: pip install openpyxl",
            file=sys.stderr,
        )
        sys.exit(1)

    excel_path = Path(args.excel_file)
    if not excel_path.exists():
        print(f"Error: no such file: {excel_path}", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else paths.PROFILE / "salary_data.json"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    print(f"Reading: {excel_path}")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    try:
        companies = []
        for sheet_name in wb.sheetnames:
            print(f"  sheet: {sheet_name}")
            companies.extend(parse_sheet(wb[sheet_name]))
    finally:
        wb.close()

    if not companies:
        print("Error: nothing could be read out of that file.", file=sys.stderr)
        print(
            "It needs a header row in the first ten rows, with a column naming "
            "the company -- 'Company', 'Empresa', 'Firma' and the rest of "
            "COMPANY_PATTERNS at the top of this script.",
            file=sys.stderr,
        )
        sys.exit(1)

    output = {
        "metadata": {
            "source": args.source or excel_path.stem,
            "index_baseline": args.baseline,
            "index_label": "Index",
            "baseline_description": args.baseline_desc or f"Index {args.baseline} = baseline",
        },
        "companies": companies,
    }

    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(output, f, ensure_ascii=False, indent=2)

    print(f"\nDone. {len(companies)} companies written to {output_path}")


if __name__ == "__main__":
    from console import use_utf8
    use_utf8()
    main()
